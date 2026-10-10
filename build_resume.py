# -*- coding: utf-8 -*-
"""生成简历改写稿 DOCX。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "简历_汉语言文学转AI_改写稿.docx"
CN = "微软雅黑"

doc = Document()

# 全局字体（含东亚字体，否则中文会回退）
st = doc.styles["Normal"]
st.font.name = CN
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.15

for sname in ("Heading 1", "List Bullet"):
    s = doc.styles[sname]
    s.font.name = CN
    s.element.rPr.rFonts.set(qn("w:eastAsia"), CN)
doc.styles["Heading 1"].font.size = Pt(12.5)
doc.styles["Heading 1"].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)


def line(text, bold=False, size=10.5, gray=False, align=None, space_after=3):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = CN
    r._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    r.bold = bold
    r.font.size = Pt(size)
    if gray:
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(lead, text):
    """加粗小标题 + 正文的条目。"""
    p = doc.add_paragraph(style="List Bullet")
    r1 = p.add_run(lead + "：")
    r1.bold = True
    r1.font.name = CN
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    r2 = p.add_run(text)
    r2.font.name = CN
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), CN)
    for r in (r1, r2):
        r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(3)
    return p


# ---------- 抬头 ----------
line("【姓名】", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
line("求职意向：AI 教研 / AI 应用（Agent） / 大模型数据与评测　｜　期望城市：北京 / 上海 / 杭州 / 深圳　｜　随时到岗",
     gray=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
line("电话：【填写】　｜　邮箱：【填写】　｜　作品集：【Coze 公开版链接】 · 【GitHub 链接】",
     gray=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

# ---------- 教育背景 ----------
doc.add_heading("教育背景", level=1)
line("【学校名称】　汉语言文学（本科）", bold=True, space_after=1)
line("【20XX.09–20XX.06】　｜　应届生身份（未缴社保），可参加校招补录与春招", gray=True, space_after=1)
line("主修课程：中国古代文学史、古代汉语、中国古典文献学、文学理论、中国现当代文学", space_after=4)

# ---------- 项目经历 ----------
doc.add_heading("项目经历", level=1)
line("古代文学考研 AI 助手（独立开发）　｜　Dify Agent + 工作流编排", bold=True, space_after=1)
line("【2025.09–至今】　｜　作品集：Coze 公开版（可访问）· GitHub（DSL 与评测集）", gray=True, space_after=3)
bullet("场景与问题",
       "古文考研备考中「出题 → 判卷 → 讲评」全流程依赖人工，评分标准模糊、反馈周期长；"
       "通用大模型在古文场景易编造出处、混淆朝代，学生无法判断答案是否可信。")
bullet("方案与实现",
       "本地 Docker Compose 部署 Dify，搭建「知识库检索 → LLM 出题 → 结构化输出」工作流并发布为工具供 Agent 调用；"
       "Agent 侧设计量化评分 Rubric（采分点拆分：名词解释 3+2+3+2、简答 2+4+2+2、论述 4+8+4+4，满分 10/10/20），"
       "配置固定「题干 / 参考答案 / 答题要点」三段输出模板、few-shot 判卷示例，"
       "并设置反幻觉约束（不自行编题、超范围礼貌拒答、引文存疑标注「原文待核对」）。")
bullet("问题定位与迭代",
       "以 8 个朝代抽样实测，发现出题朝代命中率仅 50%、单题最多触发 12 次工具调用；"
       "据此定位工作流未按朝代过滤的问题，在 Agent 侧设计并发采样（4 路）+ 朝代校验 + 拒答兜底策略，"
       "将「输出错误朝代题目」收敛为如实拒答（编题数 0），并把工作流入参过滤列为下一版改造项。")
bullet("质量与成本量化",
       "实测单轮调用耗时 13–60 秒、prompt 约 1.5 万 tokens，据此提出上下文精简与工作流改造方案；"
       "建立三段结构输出的稳定性校验机制，保证题型归一化（名词解释 10 分 / 简答 10 分 / 论述 20 分）与题干不含答案泄露。")
bullet("结果",
       "实现「出题 → 判卷 → 讲评」全流程自动闭环，并在 Coze 发布公开版；"
       "沉淀出可复用的「场景拆解 → 规则约束 → 实测验证 → 兜底迭代」的 AI 应用落地方法。")

# ---------- 实习经历 ----------
doc.add_heading("实习经历", level=1)

line("新东方教育科技集团　｜　助教", bold=True, space_after=1)
line("2025.06–2025.08", gray=True, space_after=2)
p = doc.add_paragraph(style="List Bullet")
p.add_run("负责 1 个班级 20+ 名学生的班级运营与作业批改，100% 完成学习情况跟踪，"
          "建立标准化的学习数据记录与反馈流程；").font.name = CN
p = doc.add_paragraph(style="List Bullet")
p.add_run("与主讲教师同步课堂学情、定位共性问题并配合调整辅导节奏，沉淀出可复用的学情反馈机制；"
          "在统一标准下完成大量重复性判分工作，形成对「评分尺度一致性」的实践认知。").font.name = CN
for para in doc.paragraphs[-2:]:
    para.paragraph_format.space_after = Pt(3)

line("杭州会思培训学校有限公司　｜　助教", bold=True, space_after=1)
line("2025.02–2025.03", gray=True, space_after=2)
p = doc.add_paragraph(style="List Bullet")
p.add_run("收集、分类与整理 50+ 份教学素材，搭建结构化教学素材库，形成「检索 → 归类 → 复用」的内容管理流程；").font.name = CN
p = doc.add_paragraph(style="List Bullet")
p.add_run("参与 3 个班级 40+ 名学生的学情反馈与作业批改，按统一口径输出标准化反馈记录，"
          "保证大规模重复作业下的判分一致性。").font.name = CN
for para in doc.paragraphs[-2:]:
    para.paragraph_format.space_after = Pt(3)

line("杭州帆米励托育有限公司　｜　教师", bold=True, space_after=1)
line("2024.06–2024.08", gray=True, space_after=2)
p = doc.add_paragraph(style="List Bullet")
p.add_run("对接学员与家长群体，负责日常答疑、进度跟进与突发问题处理，沟通响应获家长广泛好评；").font.name = CN
p = doc.add_paragraph(style="List Bullet")
p.add_run("独立完成 20+ 节教学视频的拍摄与剪辑，累计产出 100+ 分钟结构化教学内容，"
          "将教学知识转化为可复用的标准化内容产品。").font.name = CN
for para in doc.paragraphs[-2:]:
    para.paragraph_format.space_after = Pt(3)

# ---------- 技能与工具 ----------
doc.add_heading("技能与工具", level=1)
bullet("AI 应用", "Dify（Agent / 工作流编排 / 工具调用 / 知识库配置）、Coze、Docker Compose 本地部署")
bullet("Prompt 工程", "结构化输出模板设计、量化 Rubric 与采分点拆分、few-shot 示例、反幻觉与拒答约束")
bullet("内容与教研", "知识体系结构化、题库与素材库搭建、评分标准与 SOP 设计、古典文献检索与考据")
bullet("专业基础", "中国古代文学史、古代汉语、中国古典文献学")
bullet("语言", "英语【CET-4 / CET-6 按实际填写】，可阅读英文技术文档")

doc.save(OUT)
print("saved:", OUT)
