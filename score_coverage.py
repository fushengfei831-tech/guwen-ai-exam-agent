# -*- coding: utf-8 -*-
"""
A/B 结果评分器：
  用每道题的【答题要点】作为评分基准，计算 A/B 两组回答的「要点覆盖率」。
  方法：把每条要点转成字符二元组（bigram），统计回答中命中的比例；
        短要点（≤3 字）改用直接包含判断。
  同时统计拒答率、是否引用题库原文、分时期 / 分题型差异。
"""
import json
import re
import statistics
from collections import defaultdict

from openpyxl import load_workbook

SRC_JSON = "ab_results.json"
SRC_XLSX = "评测结果_已填充.xlsx"
OUT_XLSX = "评测结果_已评分.xlsx"
OUT_REPORT = "评测报告_要点覆盖率.txt"

PUNCT = re.compile(r"[\s【】\[\]（）()「」『』、，。；：·\-—…《》<>\"'？！,.;:!?|｜/\\]+")
THINK_LEFT = re.compile(r"</?think>|dify-deepseek-reasoning")


def norm(s):
    return PUNCT.sub("", s or "")


def bigrams(s):
    return set(s[i:i + 2] for i in range(len(s) - 1))


def parse_points(text):
    out = []
    for line in (text or "").split("\n"):
        line = line.strip()
        if line.startswith("·"):
            line = line[1:].strip()
        if line:
            out.append(line)
    return out


def core(point):
    """去掉「定义：」「性质：」这类标签前缀，取信息主体。
    题库的要点格式统一是「标签：内容」，因此一律取第一个冒号之后的部分。"""
    if "：" in point:
        head, tail = point.split("：", 1)
        if tail.strip():
            return tail.strip()
    return point


def covered(point, answer, thr=0.5):
    p = norm(core(point))
    a = norm(answer)
    if not p:
        return False
    if len(p) <= 3:
        return p in a
    bg = bigrams(p)
    if not bg:
        return p in a
    hit = sum(1 for g in bg if g in a)
    return (hit / len(bg)) >= thr


REFUSAL = re.compile(r"无法(直接)?(读取|访问|回答|确认)|没有(可引用的)?(固定)?资料|抱歉|不(能|敢)(凭空|照原文)|以免编造|不知道")
QUOTE = re.compile(r"【参考答案】|【答题要点】|【题干】")


def main():
    data = json.load(open(SRC_JSON, encoding="utf-8"))
    print("读取 %d 条实验记录" % len(data))

    leftover = sum(1 for d in data
                   for g in ("A", "B") if THINK_LEFT.search(d[g]["final"]))
    print("残留 think 标记的回答数：%d" % leftover)

    recs = []
    for d in data:
        pts = parse_points(d["points"])
        row = {"no": d["no"], "period": d["period"], "type": d["type"],
               "stem": d["stem"], "n_points": len(pts),
               "A_prompt": (d["A"]["usage"] or {}).get("prompt_tokens"),
               "B_prompt": (d["B"]["usage"] or {}).get("prompt_tokens"),
               "A_ret": d["A"]["n_retrieved"], "B_ret": d["B"]["n_retrieved"],
               "A_len": len(d["A"]["final"]), "B_len": len(d["B"]["final"])}
        for g in ("A", "B"):
            ans = d[g]["final"]
            covs = [covered(p, ans) for p in pts] if pts else []
            row[g + "_cov"] = (sum(covs) / len(covs)) if covs else 0.0
            row[g + "_full"] = bool(covs) and all(covs)
            row[g + "_refuse"] = bool(REFUSAL.search(ans))
            row[g + "_quote"] = bool(QUOTE.search(ans))
            row[g + "_detail"] = "".join("1" if c else "0" for c in covs)
        recs.append(row)

    def avg(key, sub=None):
        vals = [r[key] for r in recs if sub is None or sub(r)]
        return sum(vals) / len(vals) if vals else 0.0

    lines = []
    add = lines.append
    add("=" * 66)
    add(" 古文 AI 助手 A/B 对照实验结果（评测基准：你自己的 276 题题库）")
    add("=" * 66)
    add("")
    add("样本：%d 道题（名词解释 %d / 简答 %d），覆盖 %d 个文学史时期"
        % (len(recs),
           sum(1 for r in recs if r["type"] == "名词解释"),
           sum(1 for r in recs if r["type"] == "简答题"),
           len(set(r["period"] for r in recs))))
    add("")
    add("--- 一、总体结果 ---")
    add("指标                     A组(无知识库)     B组(带知识库)     差值")
    a_cov, b_cov = avg("A_cov"), avg("B_cov")
    add("要点覆盖率               %13.1f%% %15.1f%% %+13.1f 个百分点"
        % (a_cov * 100, b_cov * 100, (b_cov - a_cov) * 100))
    a_full = sum(1 for r in recs if r["A_full"])
    b_full = sum(1 for r in recs if r["B_full"])
    add("要点全覆盖题数           %13d  %15d %+13d" % (a_full, b_full, b_full - a_full))
    a_ref = sum(1 for r in recs if r["A_refuse"])
    b_ref = sum(1 for r in recs if r["B_refuse"])
    add("拒答/说不知道的题数      %13d  %15d %+13d" % (a_ref, b_ref, b_ref - a_ref))
    add("平均 prompt tokens       %13.0f  %15.0f %+13.0f"
        % (avg("A_prompt"), avg("B_prompt"), avg("B_prompt") - avg("A_prompt")))
    add("平均引用知识库片段数     %13.1f  %15.1f %+13.1f"
        % (avg("A_ret"), avg("B_ret"), avg("B_ret") - avg("A_ret")))
    add("平均回答字数             %13.0f  %15.0f %+13.0f"
        % (avg("A_len"), avg("B_len"), avg("B_len") - avg("A_len")))
    add("引用题库原文的题数       %13d  %15d"
        % (sum(1 for r in recs if r["A_quote"]), sum(1 for r in recs if r["B_quote"])))
    add("")
    add("--- 二、分时期对比（假设：冷门时期增益更大）---")
    add("时期            题数   A组覆盖率   B组覆盖率   差值")
    by_period = defaultdict(list)
    for r in recs:
        by_period[r["period"]].append(r)
    for p, rs in sorted(by_period.items(), key=lambda x: -len(x[1])):
        a = sum(x["A_cov"] for x in rs) / len(rs)
        b = sum(x["B_cov"] for x in rs) / len(rs)
        add("%-12s %5d %10.1f%% %10.1f%% %+9.1f" % (p, len(rs), a * 100, b * 100, (b - a) * 100))
    add("")
    add("--- 三、分题型对比 ---")
    add("题型          题数   A组覆盖率   B组覆盖率   差值")
    by_type = defaultdict(list)
    for r in recs:
        by_type[r["type"]].append(r)
    for t, rs in sorted(by_type.items(), key=lambda x: -len(x[1])):
        a = sum(x["A_cov"] for x in rs) / len(rs)
        b = sum(x["B_cov"] for x in rs) / len(rs)
        add("%-12s %5d %10.1f%% %10.1f%% %+9.1f" % (t, len(rs), a * 100, b * 100, (b - a) * 100))
    add("")
    add("--- 四、逐题明细（0/1 表示该条要点是否被覆盖）---")
    add("序 时期        题型      要点数  A覆盖  A明细  B覆盖  B明细  题干")
    for r in recs:
        add("%2d %-11s %-8s %5d %6.0f%%  %-6s %6.0f%%  %-6s %s"
            % (r["no"], r["period"], r["type"], r["n_points"],
               r["A_cov"] * 100, r["A_detail"] or "-",
               r["B_cov"] * 100, r["B_detail"] or "-",
               r["stem"][:26]))
    add("")
    add("--- 五、B 组未拿到满分的题（可深挖的失败案例）---")
    for r in recs:
        if r["B_cov"] < 1.0:
            add("  %2d. [%s] %s  覆盖率 %.0f%%  明细 %s"
                % (r["no"], r["period"], r["stem"][:30], r["B_cov"] * 100, r["B_detail"]))
    add("")
    add("--- 六、方法说明（面试可直接讲）---")
    add("1. 评测基准来自自建题库的 276 道题，按 8 个文学史时期分层抽样 30 道。")
    add("2. 评分标准用题库原有的【答题要点】，不用主观打分：把要点转成字符二元组，")
    add("   统计回答中的命中比例，>=50% 记为覆盖该要点；短要点用直接包含判断。")
    add("3. 唯一变量是「是否挂载知识库」，两个应用使用同一模型、同一提示词。")
    add("4. 局限（必须如实说明）：知识库中包含参考答案，因此本实验测量的是")
    add("   「能否检索并完整复述标准内容」，不等于独立推理能力。")

    report = "\n".join(lines)
    open(OUT_REPORT, "w", encoding="utf-8").write(report)
    print(report)

    # 写入工作簿
    wb = load_workbook(SRC_XLSX)
    ws = wb["评测题目"]
    for r in recs:
        row = r["no"] + 1
        ws.cell(row, 7, r["A_cov"])
        ws.cell(row, 7).number_format = "0%"
        ws.cell(row, 9, r["B_cov"])
        ws.cell(row, 9).number_format = "0%"
        note = "B引用了题库原文" if r["B_quote"] else ""
        if r["A_refuse"]:
            note = (note + " / A拒答").strip(" /")
        ws.cell(row, 10, note)
    wb.save(OUT_XLSX)
    print("\n已输出：%s / %s" % (OUT_REPORT, OUT_XLSX))


if __name__ == "__main__":
    main()
