# -*- coding: utf-8 -*-
"""
基于你自己的题库（完美版.txt）做两件事：
  1. 输出结构化数据 题库_结构化.csv（可直接用于知识库 / 后续处理）
  2. 分层抽样生成 评测集_你的题库_30题.xlsx（覆盖 8 个文学史时期，含评分基准）
"""
import csv
import os
import random
import re
import collections

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"D:\Users\HUAWEI\Desktop\完美版.txt"
CSV_OUT = "题库_结构化.csv"
XLSX_OUT = "评测集_你的题库_30题.xlsx"
TARGET = 30
random.seed(42)

CN = "微软雅黑"
TITLE = Font(name=CN, bold=True, size=14, color="1F3864")
H1 = Font(name=CN, bold=True, size=10.5, color="FFFFFF")
H1F = PatternFill("solid", fgColor="2F5597")
BODY = Font(name=CN, size=10)
BOLD = Font(name=CN, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEL = PatternFill("solid", fgColor="FFF2CC")
BLUE = PatternFill("solid", fgColor="E9EEF4")

# ---------- 1. 解析 ----------
txt = open(SRC, encoding="utf-8", errors="replace").read()
blocks = re.split(r"\n##\s+", txt)[1:]

recs = []
for i, b in enumerate(blocks, 1):
    def f(key):
        m = re.search(r"- %s：(.+)" % key, b)
        return m.group(1).strip() if m else ""
    m = re.search(r"【题干】\s*\n\s*(.+?)\s*\n", b, re.S)
    stem = m.group(1).strip() if m else ""
    m = re.search(r"【参考答案】\s*\n(.*?)(?=\n【答题要点】|\Z)", b, re.S)
    answer = m.group(1).strip() if m else ""
    m = re.search(r"【答题要点】\s*\n(.*?)(?=\n---|\Z)", b, re.S)
    points = [x.strip()[1:].strip() for x in (m.group(1).split("\n") if m else [])
              if x.strip().startswith("-")]
    recs.append({
        "no": i,
        "subject": f("科目"),
        "period": f("时期"),
        "type": f("题型"),
        "stem": stem,
        "answer": answer,
        "points": points,
    })

print("解析题目：%d 道" % len(recs))

# ---------- 2. 输出结构化 CSV ----------
with open(CSV_OUT, "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["题号", "科目", "时期", "题型", "题干", "参考答案", "答题要点", "要点条数"])
    for r in recs:
        w.writerow([r["no"], r["subject"], r["period"], r["type"], r["stem"],
                    r["answer"], " | ".join(r["points"]), len(r["points"])])
print("已输出：%s" % CSV_OUT)

# ---------- 3. 分层抽样 ----------
# 用名词解释 + 简答题（答案短、要点清晰，适合自动算覆盖率）
pool = [r for r in recs if r["type"] in ("名词解释", "简答题")]
by_period = collections.defaultdict(list)
for r in pool:
    by_period[r["period"]].append(r)

periods = sorted(by_period, key=lambda p: -len(by_period[p]))
total_pool = len(pool)

# 按比例分配，每期至少 2 题，不超过该期题量
alloc = {}
for p in periods:
    alloc[p] = max(2, int(round(len(by_period[p]) / total_pool * TARGET)))
for p in periods:
    alloc[p] = min(alloc[p], len(by_period[p]))

# 调整到 TARGET
while sum(alloc.values()) > TARGET:
    p = max((x for x in periods if alloc[x] > 2), key=lambda x: alloc[x], default=None)
    if p is None:
        p = max(alloc, key=lambda x: alloc[x])
    alloc[p] -= 1
while sum(alloc.values()) < TARGET:
    cand = [x for x in periods if alloc[x] < len(by_period[x])]
    if not cand:
        break
    p = max(cand, key=lambda x: len(by_period[x]) - alloc[x])
    alloc[p] += 1

print("\n分层抽样分配：")
for p in periods:
    print("  %-12s 题库 %3d 道 → 抽 %d 道" % (p, len(by_period[p]), alloc[p]))
print("  合计 %d 道" % sum(alloc.values()))

sample = []
for p in periods:
    cand = by_period[p][:]
    random.shuffle(cand)
    # 优先取答题要点 3–5 条的题（便于覆盖率判分，且信息量适中）
    cand.sort(key=lambda r: 0 if 3 <= len(r["points"]) <= 5 else 1)
    picked = []
    ni = si = 0
    for r in cand:
        if len(picked) >= alloc[p]:
            break
        # 名词解释与简答尽量交替
        if r["type"] == "名词解释" and ni > si + 1:
            continue
        if r["type"] == "简答题" and si > ni + 1:
            continue
        picked.append(r)
        if r["type"] == "名词解释":
            ni += 1
        else:
            si += 1
    for r in cand:                      # 补齐
        if len(picked) >= alloc[p]:
            break
        if r not in picked:
            picked.append(r)
    sample += picked[:alloc[p]]

sample.sort(key=lambda r: (r["period"], r["type"]))
print("\n抽样完成：%d 道（名词解释 %d / 简答 %d）"
      % (len(sample),
         sum(1 for r in sample if r["type"] == "名词解释"),
         sum(1 for r in sample if r["type"] == "简答题")))

# ---------- 4. 生成评测集工作簿 ----------
wb = Workbook()
ws = wb.active
ws.title = "评测题目"
heads = ["序号", "时期", "题型", "题干", "评分基准（答题要点）",
         "A组回答（无知识库）", "A组覆盖", "B组回答（带知识库）", "B组覆盖", "备注"]
widths = [6, 12, 11, 26, 44, 40, 9, 40, 9, 16]
for i, (h, w) in enumerate(zip(heads, widths), 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(1, i, h)
    c.font = H1; c.fill = H1F; c.alignment = CEN; c.border = BD
ws.row_dimensions[1].height = 34

for i, r in enumerate(sample, 1):
    row = i + 1
    pts = "\n".join("· " + x for x in r["points"])
    vals = [i, r["period"], r["type"], r["stem"], pts,
            None, None, None, None, None]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v)
        c.border = BD
        c.font = BOLD if j == 4 else BODY
        c.alignment = CEN if j in (1, 2, 3, 7, 9) else WRAP
    ws.cell(row, 5).fill = YEL
    ws.row_dimensions[row].height = 62

ws.freeze_panes = "D2"
last = len(sample) + 1

# 说明页
ws2 = wb.create_sheet("怎么用")
for col, w in zip("ABC", [4, 20, 96]):
    ws2.column_dimensions[col].width = w
lines = [
    ("t", "评测集（基于你自己的 276 题题库，分层抽样 30 题）", ""),
    ("b", "", ""),
    ("h", "一、为什么用你自己的题库", ""),
    ("r", "优势", "这是你的垂类资产，不是从网上抄的公共语料。面试时你能说：题库是我整理的 276 道真题，评测集是按 8 个文学史时期分层抽样出来的。"),
    ("h", "二、分层抽样设计", ""),
    ("r", "怎么抽的", "从「名词解释 + 简答题」共 154 道里抽 30 道，按 8 个时期（先秦/秦汉/魏晋南北朝/隋唐五代/宋/辽金元/明/清）题量比例分配，每个时期至少 2 道。"),
    ("r", "为什么排除论述题", "论述题回答长、判分慢，第一轮先用名词解释和简答跑通；论述题留作第二轮扩样。"),
    ("h", "三、评分基准：答题要点覆盖率（自动可算）", ""),
    ("r", "核心指标", "每道题都用你题库里的【答题要点】当评分基准。判断回答覆盖了几条要点，得到「要点覆盖率」，不需要主观打分。"),
    ("r", "为什么这个指标好", "① 标准来自你自己的真题答案，权威；② 可复现，换个模型、换个人来测结果一样；③ 机器能自动比对，30 题不用手动打分。"),
    ("h", "四、执行步骤", ""),
    ("r", "第 1 步", "把这个题库文件（完美版.txt 或我生成的 题库_结构化.csv）直接上传进你的 Dify 知识库。Dify 支持 .txt/.csv，不用转换。"),
    ("r", "第 2 步", "建两个聊天助手应用：A 组不挂知识库；B 组挂上题库知识库。系统提示词保持一致。"),
    ("r", "第 3 步", "把本表 30 道题的题干，分别发给 A 组和 B 组，把回答粘回 F 列和 H 列。"),
    ("r", "第 4 步", "把两个应用的 API 密钥给我，我用脚本自动算每条回答覆盖了几条要点，并出最终对比表。"),
    ("h", "五、预期结论（假设）", ""),
    ("r", "假设", "知识库在「冷门时期」（辽金元、明代、秦汉、清代）的增益应该最大；在「常见作家作品」（李白、杜甫、苏轼）上两组差距可能很小。"),
    ("r", "为什么这样预期", "通用大模型对最著名的作家作品背得很熟，但对冷门时期、具体作品的细节容易含糊——知识库恰好补的就是这块。"),
]
r = 1
for kind, a, b in lines:
    if kind == "t":
        ws2.cell(r, 2, a).font = TITLE; r += 1
    elif kind == "b":
        r += 1
    elif kind == "h":
        c = ws2.cell(r, 2, a); c.font = Font(name=CN, bold=True, size=11, color="1F3864")
        c.fill = PatternFill("solid", fgColor="D9E2F3")
        ws2.cell(r, 3).fill = PatternFill("solid", fgColor="D9E2F3")
        r += 1
    else:
        ca = ws2.cell(r, 2, a); ca.font = BOLD; ca.alignment = WRAP; ca.border = BD
        cb = ws2.cell(r, 3, b); cb.font = BODY; cb.alignment = WRAP; cb.border = BD
        ws2.row_dimensions[r].height = max(28, 15 * (len(b) // 52 + 1))
        r += 1

wb.save(XLSX_OUT)
print("已输出：%s" % XLSX_OUT)
print("\n抽样清单：")
for i, r in enumerate(sample, 1):
    print("  %2d. [%s/%s] %s（要点 %d 条）"
          % (i, r["period"], r["type"], r["stem"][:34], len(r["points"])))
