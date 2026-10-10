# -*- coding: utf-8 -*-
"""
A/B 对照实验执行器：
  读取 评测集_你的题库_30题.xlsx 的 30 道题
  分别调用 A 组（无知识库）与 B 组（有知识库）
  剥离 <think> 思考过程，只保留最终答案
  把回答写回新的工作簿，并把完整数据存 JSON 供评分脚本使用
"""
import json
import os
import re
import time
import urllib.error
import urllib.request

from openpyxl import load_workbook

BASE = "http://localhost/v1"
KEY_A = os.environ["KEY_A"]
KEY_B = os.environ["KEY_B"]
SRC = "评测集_你的题库_30题.xlsx"
OUT_XLSX = "评测结果_已填充.xlsx"
OUT_JSON = "ab_results.json"

THINK = re.compile(r"<think>.*?</think>", re.S)


def strip_think(t):
    t = THINK.sub("", t)
    t = re.sub(r"<!--dify-deepseek-reasoning-->", "", t)
    return t.strip()


def chat(key, query, timeout=300, retries=2):
    payload = {"query": query, "inputs": {}, "response_mode": "streaming",
               "user": "ab-eval"}
    last = None
    for attempt in range(retries + 1):
        try:
            req = urllib.request.Request(
                BASE + "/chat-messages",
                data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                headers={"Authorization": "Bearer " + key,
                         "Content-Type": "application/json"},
                method="POST")
            raw, usage, nret, t0 = "", None, 0, time.time()
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                for line in resp:
                    s = line.decode("utf-8", "replace").strip()
                    if not s.startswith("data:"):
                        continue
                    b = s[5:].strip()
                    if not b or b == "[DONE]":
                        continue
                    try:
                        ev = json.loads(b)
                    except Exception:
                        continue
                    if ev.get("event") in ("message", "agent_message"):
                        raw += ev.get("answer", "") or ""
                    elif ev.get("event") == "message_end":
                        md = ev.get("metadata", {}) or {}
                        usage = md.get("usage")
                        nret = len(md.get("retriever_resources", []) or [])
            return {"raw": raw, "final": strip_think(raw), "usage": usage,
                    "n_retrieved": nret, "elapsed": round(time.time() - t0, 1)}
        except Exception as e:
            last = e
            time.sleep(2)
    return {"raw": "", "final": "", "usage": None, "n_retrieved": 0,
            "elapsed": 0, "error": repr(last)}


def make_query(qtype, stem):
    stem = stem.strip()
    if qtype == "名词解释":
        return "请解释「%s」。" % stem
    return stem


def main():
    wb = load_workbook(SRC)
    ws = wb["评测题目"]
    rows = []
    for r in range(2, ws.max_row + 1):
        no = ws.cell(r, 1).value
        if no is None:
            continue
        rows.append({
            "row": r,
            "no": no,
            "period": ws.cell(r, 2).value,
            "type": ws.cell(r, 3).value,
            "stem": ws.cell(r, 4).value,
            "points": ws.cell(r, 5).value,
        })
    print("读取题目 %d 道" % len(rows))

    results = []
    for i, it in enumerate(rows, 1):
        q = make_query(it["type"], it["stem"])
        ra = chat(KEY_A, q)
        rb = chat(KEY_B, q)
        results.append({
            "no": it["no"], "period": it["period"], "type": it["type"],
            "stem": it["stem"], "points": it["points"], "query": q,
            "A": ra, "B": rb,
        })
        ws.cell(it["row"], 6, ra["final"])
        ws.cell(it["row"], 8, rb["final"])
        print("[%2d/%d] %s | A: %ss %s | B: %ss %s"
              % (i, len(rows), it["stem"][:22],
                 ra["elapsed"], "OK" if ra["final"] else "空",
                 rb["elapsed"], "OK" if rb["final"] else "空"), flush=True)

    wb.save(OUT_XLSX)
    json.dump(results, open(OUT_JSON, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print("\n完成：%s / %s" % (OUT_XLSX, OUT_JSON))


if __name__ == "__main__":
    main()
