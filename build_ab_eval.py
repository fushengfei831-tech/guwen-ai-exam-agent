# -*- coding: utf-8 -*-
"""生成《古文 AI 事实准确性对比评测集（30 题）》工作簿。
题目全部基于 诗经 / 楚辞 / 论语 语料，答案唯一、可客观判分。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "古文AI准确性对比评测集_30题.xlsx"

CN = "微软雅黑"
TITLE = Font(name=CN, bold=True, size=14, color="1F3864")
H1 = Font(name=CN, bold=True, size=10.5, color="FFFFFF")
H1F = PatternFill("solid", fgColor="2F5597")
SUB = Font(name=CN, bold=True, size=11, color="1F3864")
SUBF = PatternFill("solid", fgColor="D9E2F3")
BODY = Font(name=CN, size=10)
BOLD = Font(name=CN, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEL = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()

# ============================================================
# Sheet 1：使用说明
# ============================================================
ws = wb.active
ws.title = "使用说明"
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 100

rows = [
    ("t", "古文 AI 事实准确性对比评测集（30 题）", ""),
    ("t2", "目的：产出一个能写进简历的数字 —— 有知识库 vs 无知识库，答对率差多少。", ""),
    ("b", "", ""),
    ("h", "一、这个实验在做什么", ""),
    ("r", "一句话", "同一个模型、同一批问题，只改一个变量（有没有挂知识库），比较答对率。这是最标准、最容易讲清楚的对照实验。"),
    ("r", "为什么用事实题", "事实题有唯一客观答案（作者、朝代、篇目归属、原文），判分不需要主观判断，你一个人一小时能打完，而且面试官无法质疑。"),
    ("b", "", ""),
    ("h", "二、建两个应用（10 分钟）", ""),
    ("r", "A 组：对照组", "Dify 里新建一个最简单的「聊天助手」应用，模型选 DeepSeek（与你现有应用一致），**不要挂任何知识库**。命名：古文测评-A组（无知识库）。"),
    ("r", "B 组：实验组", "再新建一个完全一样的「聊天助手」，**在上下文里挂上你的古文知识库**。命名：古文测评-B组（带知识库）。"),
    ("r", "提示词保持一致", "两个应用的系统提示词都留空或写同一句：「请直接回答，不确定时请说明」。**唯一变量只能是知识库。**"),
    ("b", "", ""),
    ("h", "三、跑题与记录（约 1 小时）", ""),
    ("r", "第 1 步", "把「题目」列的问题，**逐条分别发给 A 组和 B 组**（共 60 次对话）。"),
    ("r", "第 2 步", "把两组的回答分别复制到「答题记录」表的 F 列（A 组）和 H 列（B 组）。"),
    ("r", "第 3 步", "对照「标准答案」列，在 G 列和 I 列用下拉选择判定：对 / 部分对 / 错。"),
    ("r", "判定口径", "「对」= 核心答案完全正确；「部分对」= 沾边但缺关键信息；「错」= 答错、答非所问、或明确说不知道。"),
    ("b", "", ""),
    ("h", "四、你要拿到的数字（自动计算）", ""),
    ("r", "① 答对率变化", "「结果汇总」表自动算出：A 组 __% → B 组 __%，提升 __ 个百分点。"),
    ("r", "② 分类看差异", "按「诗经篇目 / 诗经原文 / 楚辞 / 论语」四类分别看，找出知识库在哪类问题上增益最大——这是面试时最出彩的细节。"),
    ("r", "③ 诚实记录失败项", "如果某类题目 B 组并没有变好（甚至更差），**照实记录**。这比全赢更可信，也是你「下一步迭代」的依据。"),
    ("b", "", ""),
    ("h", "五、写进简历的句子模板", ""),
    ("r", "模板", "构建 30 题古文事实准确性评测集（覆盖诗经篇目归属、原文默写、楚辞篇目、论语语录出处四类），采用同模型对照实验（唯一变量为是否挂载知识库），实测答对率由 __% 提升至 __%（+__ 个百分点），其中 ___ 类问题增益最大（__% → __%）。"),
    ("b", "", ""),
    ("h", "六、面试会追问的 3 个问题（先想好答案）", ""),
    ("r", "「30 题够吗？」", "答：够定位系统性问题，不够做学术结论。我关心的是「知识库对哪类问题有效」，30 题已经能看出明显差异。"),
    ("r", "「怎么保证公平？」", "答：同一个模型、同一个提示词、同一批问题，唯一变量是知识库；判定对照预先写好的标准答案，不看是哪组答的。"),
    ("r", "「有没有反例？」", "答：有，我会把反例写进结论——比如常见名句两组都能答对，说明知识库的价值集中在冷门篇目和原文精确性上。"),
]
r = 1
for kind, a, b in rows:
    if kind == "t":
        ws.cell(r, 2, a).font = TITLE
        r += 1
    elif kind == "t2":
        ws.cell(r, 2, a).font = Font(name=CN, size=10.5, italic=True, color="595959")
        r += 1
    elif kind == "b":
        r += 1
    elif kind == "h":
        c = ws.cell(r, 2, a); c.font = SUB; c.fill = SUBF
        ws.cell(r, 3).fill = SUBF
        r += 1
    else:
        ca = ws.cell(r, 2, a); ca.font = BOLD; ca.alignment = WRAP; ca.border = BD
        cb = ws.cell(r, 3, b); cb.font = BODY; cb.alignment = WRAP; cb.border = BD
        ws.row_dimensions[r].height = max(28, 15 * (len(b) // 55 + 1))
        r += 1

# ============================================================
# Sheet 2：答题记录
# ============================================================
ws2 = wb.create_sheet("答题记录")
heads = ["题号", "类别", "题目", "标准答案", "出处",
         "A组回答（无知识库）", "A判定", "B组回答（带知识库）", "B判定", "备注"]
widths = [6, 13, 34, 30, 18, 34, 9, 34, 9, 22]
for i, (h, w) in enumerate(zip(heads, widths), 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
    c = ws2.cell(1, i, h)
    c.font = H1; c.fill = H1F; c.alignment = CEN; c.border = BD
ws2.row_dimensions[1].height = 32

Q = [
    # ---- 诗经·篇目归属 ----
    ("诗经篇目", "《诗经》按音乐性质分为哪三部分？", "风、雅、颂", "《诗经》体例"),
    ("诗经篇目", "《诗经》共收录多少篇作品？", "305 篇（另有 6 篇笙诗有目无辞）", "《诗经》总论"),
    ("诗经篇目", "《诗经》中的「颂」包括哪三部分？", "周颂、鲁颂、商颂", "《诗经》体例"),
    ("诗经篇目", "《诗经》十五国风中的「王风」指哪个地区的诗？", "东周王畿（今洛阳一带）的诗", "《王风》"),
    ("诗经篇目", "「鸤鸠」一诗属于十五国风中的哪一国？", "曹风（《曹风·鸤鸠》）", "《曹风·鸤鸠》"),
    ("诗经篇目", "「硕鼠」一诗属于十五国风中的哪一国？", "魏风（《魏风·硕鼠》）", "《魏风·硕鼠》"),
    ("诗经篇目", "「七月」一诗属于十五国风中的哪一国？", "豳风（《豳风·七月》）", "《豳风·七月》"),
    # ---- 诗经·原文 ----
    ("诗经原文", "「关关雎鸠，在河之洲」出自哪一篇？", "《周南·关雎》", "《周南·关雎》"),
    ("诗经原文", "「蒹葭苍苍，白露为霜」出自哪一篇？", "《秦风·蒹葭》", "《秦风·蒹葭》"),
    ("诗经原文", "「昔我往矣，杨柳依依」出自哪一篇？", "《小雅·采薇》", "《小雅·采薇》"),
    ("诗经原文", "「投我以木桃，报之以琼瑶」出自哪一篇？", "《卫风·木瓜》", "《卫风·木瓜》"),
    ("诗经原文", "「他山之石，可以攻玉」出自哪一篇？", "《小雅·鹤鸣》", "《小雅·鹤鸣》"),
    ("诗经原文", "「呦呦鹿鸣，食野之苹」出自哪一篇？", "《小雅·鹿鸣》", "《小雅·鹿鸣》"),
    ("诗经原文", "《卫风·氓》中「桑之未落」的下一句是什么？", "其叶沃若", "《卫风·氓》"),
    # ---- 楚辞 ----
    ("楚辞", "《楚辞》中《九歌》共多少篇？", "11 篇", "《九歌》"),
    ("楚辞", "《楚辞》中《九章》共多少篇？", "9 篇", "《九章》"),
    ("楚辞", "「路漫漫其修远兮，吾将上下而求索」出自哪一篇？", "《离骚》", "《离骚》"),
    ("楚辞", "「袅袅兮秋风，洞庭波兮木叶下」出自哪一篇？", "《九歌·湘夫人》", "《九歌·湘夫人》"),
    ("楚辞", "「举世皆浊我独清，众人皆醉我独醒」出自哪一篇？", "《渔父》", "《渔父》"),
    ("楚辞", "「沧浪之水清兮，可以濯吾缨」出自哪一篇？", "《渔父》", "《渔父》"),
    ("楚辞", "《楚辞》中《天问》的作者是谁？", "屈原", "《天问》"),
    ("楚辞", "《离骚》开篇两句是什么？", "帝高阳之苗裔兮，朕皇考曰伯庸。", "《离骚》"),
    # ---- 论语 ----
    ("论语", "《论语》共多少篇？", "20 篇", "《论语》体例"),
    ("论语", "「三人行，必有我师焉」出自哪一篇？", "《述而》", "《论语·述而》"),
    ("论语", "「岁寒，然后知松柏之后凋也」出自哪一篇？", "《子罕》", "《论语·子罕》"),
    ("论语", "「知之为知之，不知为不知，是知也」出自哪一篇？", "《为政》", "《论语·为政》"),
    ("论语", "「工欲善其事，必先利其器」出自哪一篇？", "《卫灵公》", "《论语·卫灵公》"),
    ("论语", "「子在川上曰：逝者如斯夫，不舍昼夜」出自哪一篇？", "《子罕》", "《论语·子罕》"),
    ("论语", "「学而时习之，不亦说乎」出自哪一篇？", "《学而》", "《论语·学而》"),
    ("论语", "「己所不欲，勿施于人」在《论语》中出现于哪两篇？", "《颜渊》与《卫灵公》", "《论语》"),
]
assert len(Q) == 30, len(Q)

for i, (cat, q, ans, src) in enumerate(Q, 1):
    r = i + 1
    vals = [i, cat, q, ans, src, None, None, None, None, None]
    for j, v in enumerate(vals, 1):
        c = ws2.cell(r, j, v)
        c.border = BD
        c.font = BODY if j not in (4,) else BOLD
        c.alignment = CEN if j in (1, 2, 7, 9) else WRAP
    ws2.cell(r, 4).fill = YEL
    ws2.row_dimensions[r].height = 34

ws2.freeze_panes = "C2"
last = len(Q) + 1

dv = DataValidation(type="list", formula1='"对,部分对,错"', allow_blank=True)
ws2.add_data_validation(dv)
dv.add(f"G2:G{last}")
dv.add(f"I2:I{last}")

# ============================================================
# Sheet 3：结果汇总
# ============================================================
ws3 = wb.create_sheet("结果汇总")
for col, w in zip("ABCDEF", [4, 22, 16, 16, 16, 40]):
    ws3.column_dimensions[col].width = w
ws3["B1"] = "结果汇总（自动计算，填完判定即出结果）"
ws3["B1"].font = TITLE

hdr = ["指标", "A组（无知识库）", "B组（带知识库）", "差值", "备注"]
for i, h in enumerate(hdr, 2):
    c = ws3.cell(3, i, h)
    c.font = H1; c.fill = H1F; c.alignment = CEN; c.border = BD

J = f"答题记录!$G$2:$G${last}"   # A 判定
K = f"答题记录!$I$2:$I${last}"   # B 判定
CAT = f"答题记录!$B$2:$B${last}"
N = len(Q)

items = [
    ("答对题数", f'=COUNTIF({J},"对")', f'=COUNTIF({K},"对")', "核心数字"),
    ("部分对题数", f'=COUNTIF({J},"部分对")', f'=COUNTIF({K},"部分对")', ""),
    ("答错题数", f'=COUNTIF({J},"错")', f'=COUNTIF({K},"错")', ""),
    ("未判定题数", f'=COUNTBLANK({J})', f'=COUNTBLANK({K})', "填完判定后应变为 0"),
]
r = 4
for name, fa, fb, note in items:
    ws3.cell(r, 2, name).font = BOLD
    ws3.cell(r, 3, fa)
    ws3.cell(r, 4, fb)
    ws3.cell(r, 5, f"=IFERROR(D{r}-C{r},0)")
    ws3.cell(r, 6, note).font = Font(name=CN, size=9, color="595959")
    for j in range(2, 7):
        ws3.cell(r, j).border = BD
        if j in (3, 4, 5):
            ws3.cell(r, j).alignment = CEN
    r += 1

ws3.cell(r, 2, "答对率").font = BOLD
ws3.cell(r, 3, f"=IFERROR(C4/{N},0)")
ws3.cell(r, 4, f"=IFERROR(D4/{N},0)")
ws3.cell(r, 5, f"=IFERROR(D{r}-C{r},0)")
ws3.cell(r, 6, "★ 写进简历的数字").font = Font(name=CN, size=9, color="C00000")
for j in range(2, 7):
    ws3.cell(r, j).border = BD
    if j in (3, 4, 5):
        ws3.cell(r, j).alignment = CEN
        ws3.cell(r, j).number_format = "0.0%"
        ws3.cell(r, j).fill = YEL
rate_row = r
r += 2

# 分类汇总
ws3.cell(r, 2, "按类别对比（面试时的细节亮点）").font = SUB
ws3.cell(r, 2).fill = SUBF
for j in (3, 4, 5, 6):
    ws3.cell(r, j).fill = SUBF
r += 1
for i, h in enumerate(["类别", "A组答对", "B组答对", "差值", "该类题数"], 2):
    c = ws3.cell(r, i, h)
    c.font = H1; c.fill = H1F; c.alignment = CEN; c.border = BD
r += 1
cats = ["诗经篇目", "诗经原文", "楚辞", "论语"]
for cat in cats:
    ws3.cell(r, 2, cat).font = BODY
    ws3.cell(r, 3, f'=COUNTIFS({CAT},"{cat}",{J},"对")')
    ws3.cell(r, 4, f'=COUNTIFS({CAT},"{cat}",{K},"对")')
    ws3.cell(r, 5, f"=IFERROR(D{r}-C{r},0)")
    ws3.cell(r, 6, f'=COUNTIFS({CAT},"{cat}")')
    for j in range(2, 7):
        ws3.cell(r, j).border = BD
        if j > 2:
            ws3.cell(r, j).alignment = CEN
    r += 1
r += 1

ws3.cell(r, 2, "结论区（手动填写，直接用于简历与作品集）").font = SUB
ws3.cell(r, 2).fill = SUBF
for j in (3, 4, 5, 6):
    ws3.cell(r, j).fill = SUBF
r += 1
concl = [
    "1. 答对率由 ____% 提升至 ____%，提升 ____ 个百分点。",
    "2. 增益最大的类别是 __________（____% → ____%）。",
    "3. 增益最小或为负的类别是 __________，原因：____________________。",
    "4. 典型失败案例（A 组答错、B 组答对）：____________________________。",
    "5. 典型反例（两组都错）：____________________________，下一步优化方向：______。",
]
for t in concl:
    c = ws3.cell(r, 2, t)
    c.font = BODY
    c.alignment = WRAP
    ws3.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws3.row_dimensions[r].height = 22
    r += 1

wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
print("questions:", len(Q))
from collections import Counter
print("by category:", dict(Counter(x[0] for x in Q)))
